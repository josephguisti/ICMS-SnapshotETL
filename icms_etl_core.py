"""
icms_etl_core.py
~~~~~~~~~~~~~~~~
Core ETL logic for the ICMS weekly snapshot pipeline.

The GUI (icms_gui.py) and CLI (update_icms_snapshot_files.py) both call
``run_snapshot_files``; neither contains import logic of its own. Field and
worksheet mappings live only in ``icms_catalogue.py``.
"""

from __future__ import annotations

import io
import math
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
import pandas as pd

import icms_catalogue as catalogue
from icms_catalogue import FieldSpec, normalize_header

# ── Field lists (derived from the catalogue) ──────────────────────────────────

METADATA_FIELDS: List[str] = [
    "snapshot_date",
    "reporting_month",
    "source_file_name",
]

ALL_APPROVED_FIELDS: List[str] = [f.output for f in catalogue.PARTICIPANT_FIELDS]
ESSENTIAL_PARTICIPANT_FIELDS: List[str] = [
    f.output for f in catalogue.PARTICIPANT_FIELDS if f.essential
]
MOST_RECENT_FIELDS: List[str] = [
    f for f in ALL_APPROVED_FIELDS if f.startswith("Most Recent ")
]
EXIT_HISTORY_FIELDS: List[str] = [f.output for f in catalogue.EXIT_FIELDS]
VACANT_FIELDS: List[str] = [f.output for f in catalogue.VACANT_FIELDS]
SERVICES_LOG_FIELDS: List[str] = [f.output for f in catalogue.SERVICES_FIELDS]

MASTER_TRACKING_FIELDS: List[str] = [
    "present_in_latest_snapshot",
    "client_status",
]

# Exit-only columns carried on the master.
EXIT_FIELDS: List[str] = [
    "Check-In Date",
    "Check-Out Date",
    "Move-In Date",
    "Move-Out Date",
    "Exit Reason",
    "Exit Destination",
    "Exit Location",
    "Transferred to New Provider Within ICMS Program?",
    "New ICMS Provider",
]

# Participant rows without a client or project cannot be placed anywhere.
ROW_KEY_FIELDS: List[str] = ["CHAMP ID", "ICMS Project Name"]
# Master grain: one row per enrollment.
KEY_FIELDS: List[str] = ["CHAMP ID", "ICMS Project Name", "ICMS Enrollment ID"]
ID_FIELDS = {"CHAMP ID", "ICMS Enrollment ID"}

MASTER_COLUMNS: List[str] = (
    METADATA_FIELDS + ALL_APPROVED_FIELDS + MASTER_TRACKING_FIELDS + EXIT_FIELDS
)
FLATFILE_COLUMNS: List[str] = (
    METADATA_FIELDS + ALL_APPROVED_FIELDS + ["is_vacant_slot", "# Days Vacant"]
)


def _unique(items: Iterable[str]) -> List[str]:
    seen: Dict[str, None] = {}
    for item in items:
        seen.setdefault(item, None)
    return list(seen)


# Calendar-date columns. Excel writers emit real date cells (not ISO text);
# CSV writers persist YYYY-MM-DD and never the pandas token "NaT".
# reporting_month (YYYY-MM) is intentionally excluded — it is not a date.
DATE_FIELDS: List[str] = _unique(
    ["snapshot_date"]
    + [
        f.output
        for specs in catalogue.FIELDS_BY_CATEGORY.values()
        for f in specs
        if f.is_date
    ]
)

_EXCEL_DATE_FORMAT = "YYYY-MM-DD"
_NULL_STRING_VALUES = {"nan": "", "None": "", "<NA>": "", "NaT": "", "NaN": ""}

DASHBOARD_FILENAME = "icms_current_dashboard_flatfile.xlsx"
SNAPSHOT_HISTORY_FILENAME = "icms_snapshot_history.csv"
EXIT_HISTORY_FILENAME = "icms_exit_log_history.csv"
SERVICES_HISTORY_FILENAME = "icms_services_log_history.csv"
MASTER_FILENAME = "icms_client_project_master.xlsx"
INGEST_LOG_FILENAME = "icms_ingest_log.csv"
DUPLICATE_WARNING_FILENAME = "icms_duplicate_client_project_warning.csv"
FORMAT_MARKER_FILENAME = "icms_output_format.txt"
BACKUP_DIRNAME = "backups"

OUTPUT_FILENAMES: List[str] = [
    DASHBOARD_FILENAME,
    SNAPSHOT_HISTORY_FILENAME,
    EXIT_HISTORY_FILENAME,
    SERVICES_HISTORY_FILENAME,
    MASTER_FILENAME,
    INGEST_LOG_FILENAME,
    DUPLICATE_WARNING_FILENAME,
    FORMAT_MARKER_FILENAME,
]

# Bump when stored outputs change shape; triggers a one-time automatic backup.
OUTPUT_FORMAT_VERSION = "2: enrollment-level master, catalogue mappings"

DEFAULT_SHEET_NAMES: Dict[str, str] = {
    cat: spec.names[0] for cat, spec in catalogue.SHEETS.items()
}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _safe_str(val) -> str:
    """Convert a cell value to a clean string; return '' for null/NaN."""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    return str(val).strip()


def _strip_internal_cols(df: pd.DataFrame) -> pd.DataFrame:
    """Drop any column whose name starts with '_' before writing outputs."""
    return df.drop(
        columns=[c for c in df.columns if str(c).startswith("_")],
        errors="ignore",
    )


def _parse_date_series(series: pd.Series) -> pd.Series:
    """Parse a column to datetime64[ns]; blanks, 'NaT', and junk become NaT."""
    return pd.to_datetime(series, errors="coerce", format="mixed")


def _format_date_columns_iso(df: pd.DataFrame) -> pd.DataFrame:
    """Persist date columns as YYYY-MM-DD strings; missing values as ''."""
    out = df.copy()
    for col in DATE_FIELDS:
        if col not in out.columns:
            continue
        parsed = _parse_date_series(out[col])
        formatted = parsed.dt.strftime("%Y-%m-%d")
        out[col] = formatted.where(parsed.notna(), "")
    return out


def _stringify_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Stringify a frame for CSV history: ISO dates, null tokens as ''."""
    out = _format_date_columns_iso(df)
    return out.astype(str).replace(_NULL_STRING_VALUES)


def _order_columns(df: pd.DataFrame, preferred: Sequence[str]) -> pd.DataFrame:
    """Catalogue columns first in catalogue order, then any legacy extras."""
    extras = [c for c in df.columns if c not in preferred]
    return df.reindex(columns=list(preferred) + extras, fill_value="")


def _write_excel(df: pd.DataFrame, path: Path) -> None:
    """Write ``df`` to ``path`` with DATE_FIELDS stored as Excel dates.

    pandas ``to_excel`` writes Python strings as Excel text. US Excel will
    not treat ISO ``YYYY-MM-DD`` text as dates, and Format Cells → Date
    cannot convert text. Parsing to date objects before write stores real
    date serials so AutoFilter date grouping works.
    """
    out = df.copy()
    date_col_names: List[str] = []
    for col in DATE_FIELDS:
        if col not in out.columns:
            continue
        parsed = _parse_date_series(out[col])
        # Python date objects (not midnight datetimes) so Excel uses a
        # date number format instead of YYYY-MM-DD HH:MM:SS.
        out[col] = [ts.date() if pd.notna(ts) else None for ts in parsed]
        date_col_names.append(col)

    with pd.ExcelWriter(
        path,
        engine="openpyxl",
        date_format=_EXCEL_DATE_FORMAT,
        datetime_format=_EXCEL_DATE_FORMAT,
    ) as writer:
        out.to_excel(writer, index=False)
        ws = writer.book.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for col_name in date_col_names:
            if col_name not in headers:
                continue
            idx = headers.index(col_name) + 1
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell = row[0]
                if cell.value is None or cell.value == "":
                    cell.value = None
                cell.number_format = _EXCEL_DATE_FORMAT


def _read_history_csv(path: Path, label: str) -> pd.DataFrame:
    """Read a stored history CSV as strings; never silently start fresh."""
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path, dtype=str, encoding="utf-8").fillna("")
    except pd.errors.EmptyDataError:
        return pd.DataFrame()
    except Exception as exc:
        raise RuntimeError(
            f"Could not read existing {label} ({path.name}): {exc}. "
            "Nothing was written. Close the file if it is open in Excel and retry."
        ) from exc


def _norm_cell(x):
    """Strip strings; turn whole-number floats into integer strings."""
    if isinstance(x, str):
        return x.strip()
    if isinstance(x, float):
        if math.isnan(x) or math.isinf(x):
            return None
        if x == int(x):
            return str(int(x))
    return x


_WHOLE_NUMBER_TEXT = re.compile(r"^(\d+)\.0+$")


def normalize_identifier(x) -> str:
    """Normalize an identifier cell to text: 100300.0, 100300, '100300 ' → '100300'."""
    if x is None:
        return ""
    if isinstance(x, bool):
        return str(x)
    if isinstance(x, (int, np.integer)):
        return str(int(x))
    if isinstance(x, (float, np.floating)):
        if math.isnan(x) or math.isinf(x):
            return ""
        return str(int(x)) if x == int(x) else str(x)
    text = _safe_str(x)
    m = _WHOLE_NUMBER_TEXT.match(text)
    return m.group(1) if m else text


def _norm_id_series(series: pd.Series) -> pd.Series:
    return series.map(normalize_identifier)


# ── Dates ─────────────────────────────────────────────────────────────────────

def parse_date_from_filename(filename: str) -> Optional[str]:
    """Search *filename* for a ``MM.DD.YYYY`` date; return ``YYYY-MM-DD``."""
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", filename)
    if not m:
        return None
    mm, dd, yyyy = m.groups()
    try:
        return date(int(yyyy), int(mm), int(dd)).isoformat()
    except ValueError:
        return None


_US_DATE = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b")
_ISO_DATE = re.compile(r"\b(\d{4})-(\d{2})-(\d{2})\b")


def _coerce_metadata_date(value) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str):
        return None
    m = _US_DATE.search(value)
    try:
        if m:
            mm, dd, yyyy = m.groups()
            return date(int(yyyy), int(mm), int(dd)).isoformat()
        m = _ISO_DATE.search(value)
        if m:
            yyyy, mm, dd = m.groups()
            return date(int(yyyy), int(mm), int(dd)).isoformat()
    except ValueError:
        return None
    return None


def read_generated_on(wb: openpyxl.Workbook) -> Optional[str]:
    """Return the workbook's ``Generated On`` date as ``YYYY-MM-DD``.

    Finds the cell labelled "Generated On" on the metadata tab and reads the
    value beneath or beside it, so extra columns (e.g. a Row ID) do not
    matter. Falls back to the only date-like value on that tab.
    """
    target = normalize_header(catalogue.METADATA_SHEET)
    sheet = next(
        (s for s in wb.sheetnames if normalize_header(s) == target), None
    )
    if sheet is None:
        return None
    rows = [list(r) for r in wb[sheet].iter_rows(max_row=50, values_only=True)]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if normalize_header(value) != target:
                continue
            below = rows[r + 1][c] if r + 1 < len(rows) and c < len(rows[r + 1]) else None
            right = row[c + 1] if c + 1 < len(row) else None
            for candidate in (below, right):
                parsed = _coerce_metadata_date(candidate)
                if parsed:
                    return parsed
    found = {
        parsed
        for row in rows
        for value in row
        if (parsed := _coerce_metadata_date(value))
    }
    return found.pop() if len(found) == 1 else None


def read_generated_on_from_file(
    path: str | Path, password: Optional[str] = ""
) -> Optional[str]:
    """Cheap ``Generated On`` lookup for date confirmation; None if unreadable."""
    try:
        wb, _ = open_workbook(path, password=password, read_only=True)
    except Exception:
        return None
    try:
        return read_generated_on(wb)
    except Exception:
        return None
    finally:
        wb.close()


def default_snapshot_date(
    path: str | Path, password: Optional[str] = ""
) -> Dict[str, Optional[str]]:
    """Return the default snapshot date for a workbook and where it came from.

    Keys: ``date``, ``source`` ("Generated On", "filename", or ""),
    ``generated_on``, ``filename_date``, ``mismatch`` (both present and differ).
    """
    generated = read_generated_on_from_file(path, password)
    from_name = parse_date_from_filename(Path(path).name)
    chosen = generated or from_name
    return {
        "date": chosen,
        "source": "Generated On" if generated else ("filename" if from_name else ""),
        "generated_on": generated,
        "filename_date": from_name,
        "mismatch": bool(generated and from_name and generated != from_name),
    }


# ── Workbook opening ──────────────────────────────────────────────────────────

def open_workbook(
    path: str | Path,
    password: Optional[str] = None,
    *,
    read_only: bool = False,
) -> Tuple[openpyxl.Workbook, bool]:
    """Open a normal or password-protected Excel workbook.

    Returns
    -------
    (workbook, was_password_protected)

    Password handling
    -----------------
    - If ``password`` is supplied it is used directly (GUI path).
    - If the file is encrypted and ``password`` is None, the user is
      prompted securely via ``getpass`` (CLI path).
    - Decryption is performed entirely in memory via msoffcrypto-tool;
      no decrypted copy is written to disk and the password is not stored.

    Raises
    ------
    FileNotFoundError
        If the input file does not exist.
    RuntimeError
        If the file cannot be opened or decrypted.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")

    permission_msg = (
        "Permission denied while opening the workbook.\n\n"
        f"File: {path}\n\n"
        "Windows blocked read access to this file. Common causes:\n"
        "- the workbook is open in Excel or another program,\n"
        "- OneDrive is syncing or has not fully downloaded the file,\n"
        "- the file is in a folder where this app does not have read permission.\n\n"
        "Close the workbook, wait for OneDrive sync to finish, or copy the file "
        "to a local folder such as Desktop and try again."
    )

    def _load(src):
        return openpyxl.load_workbook(src, data_only=True, read_only=read_only)

    # Attempt to import msoffcrypto (optional dependency)
    try:
        import msoffcrypto  # type: ignore
        _has_msoffcrypto = True
    except ImportError:
        msoffcrypto = None
        _has_msoffcrypto = False

    # ── Detect encryption ─────────────────────────────────────────────────
    is_encrypted = False
    if _has_msoffcrypto:
        try:
            with open(path, "rb") as fh:
                of = msoffcrypto.OfficeFile(fh)
                is_encrypted = of.is_encrypted()
        except PermissionError as exc:
            raise RuntimeError(permission_msg) from exc
        except Exception as exc:
            raise RuntimeError(f"Could not inspect workbook: {exc}") from exc
    else:
        # Without msoffcrypto, attempt a direct open; openpyxl will raise
        # on encrypted files and we surface a clear install message.
        try:
            return _load(path), False
        except PermissionError as exc:
            raise RuntimeError(permission_msg) from exc
        except Exception:
            raise RuntimeError(
                "The workbook could not be opened.  It may be password-protected.\n"
                "Install the decryption library with:\n"
                "    pip install msoffcrypto-tool"
            )

    # ── Plain (unencrypted) file ───────────────────────────────────────────
    if not is_encrypted:
        try:
            return _load(path), False
        except PermissionError as exc:
            raise RuntimeError(permission_msg) from exc
        except Exception as exc:
            raise RuntimeError(f"Could not open workbook: {exc}") from exc

    # ── Encrypted file ────────────────────────────────────────────────────
    if password is None:
        import getpass
        password = getpass.getpass(f"Password for '{path.name}': ")

    try:
        with open(path, "rb") as fh:
            of = msoffcrypto.OfficeFile(fh)
            of.load_key(password=password)
            buf = io.BytesIO()
            of.decrypt(buf)
        buf.seek(0)
        return _load(buf), True
    except PermissionError as exc:
        raise RuntimeError(permission_msg) from exc
    except Exception as exc:
        raise RuntimeError(
            f"Could not decrypt workbook (wrong password?): {exc}"
        ) from exc


def is_encrypted(path: str | Path) -> bool:
    """True if the workbook is password-protected (False if undeterminable)."""
    try:
        import msoffcrypto  # type: ignore
        with open(path, "rb") as fh:
            return bool(msoffcrypto.OfficeFile(fh).is_encrypted())
    except Exception:
        return False


def list_sheet_names(wb: openpyxl.Workbook) -> List[str]:
    """Return the list of sheet names in the workbook."""
    return wb.sheetnames


# ── Structural findings ───────────────────────────────────────────────────────
# A finding is a plain dict so it can be grouped, logged, and copied as text.

FINDING_LABELS: Dict[str, str] = {
    "essential_missing": "File not imported: essential column missing",
    "essential_conflict": "File not imported: conflicting identity columns",
    "category_skipped": "Tab not imported",
    "missing_sheet": "Supported tab not found (not imported)",
    "ambiguous_sheet": "Several candidate tabs differ (none imported)",
    "unfamiliar_sheet": "Unfamiliar tab (not imported)",
    "unfamiliar_header": "Unfamiliar column (not exported)",
    "missing_field": "Expected column missing (left blank)",
    "value_conflict": "Two accepted column names both present and disagree",
    "blank_enrollment": "Participant rows with blank ICMS Enrollment ID",
    "date_mismatch": "Snapshot date differs from Generated On or filename",
    "generated_on_missing": "Generated On date not found (filename date used)",
}

# Findings that mean data the analyst expected is not in the outputs.
INCOMPLETE_KINDS = {
    "essential_missing", "essential_conflict", "category_skipped",
    "missing_sheet", "ambiguous_sheet",
}


def _finding(kind: str, category: str, item: str, detail: str = "") -> Dict[str, str]:
    return {
        "kind": kind,
        "category": catalogue.CATEGORY_LABELS.get(category, category),
        "item": item,
        "detail": detail,
    }


def group_findings(per_file: Sequence[Tuple[str, Sequence[Dict[str, str]]]]) -> List[Dict]:
    """Group identical findings across files: one entry per (kind, tab, item)."""
    groups: Dict[Tuple[str, str, str], Dict] = {}
    for source, findings in per_file:
        for f in findings:
            key = (f["kind"], f["category"], f["item"])
            g = groups.setdefault(
                key, {**f, "files": [], "details": []}
            )
            if source not in g["files"]:
                g["files"].append(source)
            if f.get("detail") and f["detail"] not in g["details"]:
                g["details"].append(f["detail"])
    order = list(FINDING_LABELS)
    return sorted(
        groups.values(),
        key=lambda g: (order.index(g["kind"]) if g["kind"] in order else 99,
                       g["category"], g["item"]),
    )


def format_findings_text(groups: Sequence[Dict], total_files: int = 1) -> str:
    """Plain-text findings summary suitable for copying to the director."""
    if not groups:
        return ""
    lines: List[str] = []
    current = None
    for g in groups:
        label = FINDING_LABELS.get(g["kind"], g["kind"])
        if label != current:
            if current is not None:
                lines.append("")
            lines.append(label)
            current = label
        where = f"[{g['category']}] " if g["category"] else ""
        line = f"  - {where}{g['item']}"
        if total_files > 1:
            n = len(g["files"])
            line += f"  ({n} of {total_files} files)"
        lines.append(line)
        for d in g["details"][:3]:
            lines.append(f"      {d}")
        if total_files > 1 and len(g["files"]) < total_files:
            for name in g["files"][:5]:
                lines.append(f"      {name}")
            if len(g["files"]) > 5:
                lines.append(f"      … and {len(g['files']) - 5} more")
    return "\n".join(lines)


# ── Worksheet resolution and reading ──────────────────────────────────────────

@dataclass
class CategoryRead:
    """Result of reading one supported worksheet through the catalogue."""
    frame: pd.DataFrame
    findings: List[Dict[str, str]] = field(default_factory=list)
    essential_problems: List[str] = field(default_factory=list)
    header_row: int = 0
    sheet_name: str = ""


def _accepted_names(specs: Sequence[FieldSpec]) -> set:
    return {normalize_header(n) for s in specs for n in s.accepted}


def _detect_header_row(rows: Sequence[Sequence], accepted: set, max_scan: int = 5) -> int:
    best_idx, best_score = 0, 0
    for idx, row in enumerate(rows[:max_scan]):
        score = len({normalize_header(v) for v in (row or ()) if v is not None} & accepted)
        if score > best_score:
            best_idx, best_score = idx, score
    return best_idx


def _field_expected(spec: FieldSpec, snapshot_date: str) -> bool:
    if spec.retired or not spec.report_missing:
        return False
    return spec.expected_since is None or snapshot_date >= spec.expected_since


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def read_category(
    wb: openpyxl.Workbook,
    sheet_name: str,
    category: str,
    snapshot_date: str,
) -> CategoryRead:
    """Read a supported worksheet into catalogue output columns.

    Headers are matched by position after normalization, so two accepted
    names for one field may coexist. When they do, the first non-blank value
    in catalogue priority order is kept and any disagreement is recorded.
    Disagreement on an essential (identity) field is an essential problem:
    the caller must not import the tab rather than risk misattribution.
    """
    specs = catalogue.FIELDS_BY_CATEGORY[category]
    accepted = _accepted_names(specs)
    excluded = {normalize_header(h) for h in catalogue.KNOWN_EXCLUDED_HEADERS.get(category, ())}
    columns = [s.output for s in specs]
    result = CategoryRead(frame=pd.DataFrame(columns=columns + ["_source_row_number"]),
                          sheet_name=sheet_name)

    rows = [tuple(r) for r in wb[sheet_name].iter_rows(values_only=True)]
    if not rows:
        return result
    header_idx = _detect_header_row(rows, accepted)
    result.header_row = header_idx
    headers = list(rows[header_idx])
    norm_headers = [normalize_header(h) for h in headers]
    width = len(headers)

    data_rows = []
    row_numbers = []
    for offset, raw in enumerate(rows[header_idx + 1:]):
        padded = tuple(raw) + (None,) * max(0, width - len(raw))
        if all(_is_blank(v) for v in padded):
            continue
        data_rows.append(padded)
        row_numbers.append(header_idx + 2 + offset)

    for pos, (raw, norm) in enumerate(zip(headers, norm_headers)):
        if not norm or norm in accepted or norm in excluded:
            continue
        result.findings.append(_finding("unfamiliar_header", category, str(raw).strip()))

    data: Dict[str, List] = {}
    for spec in specs:
        positions = [
            i for name in spec.accepted for i, h in enumerate(norm_headers)
            if h == normalize_header(name)
        ]
        positions = list(dict.fromkeys(positions))
        is_id = spec.output in ID_FIELDS
        clean = normalize_identifier if is_id else _norm_cell
        if not positions:
            data[spec.output] = [None] * len(data_rows)
            if spec.essential:
                result.essential_problems.append(f"'{spec.output}' column not found")
                result.findings.append(_finding("essential_missing", category, spec.output))
            elif _field_expected(spec, snapshot_date):
                result.findings.append(_finding("missing_field", category, spec.output))
            continue

        values: List = []
        conflicts = 0
        for row in data_rows:
            present = [clean(row[i]) for i in positions]
            present = [v for v in present if not _is_blank(v)]
            values.append(present[0] if present else None)
            if len({_safe_str(v) for v in present}) > 1:
                conflicts += 1
        if is_id:
            values = [v if v else None for v in values]
        data[spec.output] = values

        if len(positions) > 1 and conflicts:
            names = " / ".join(str(headers[i]).strip() for i in positions)
            detail = (f"{names}: {conflicts} row(s) disagree; kept "
                      f"'{str(headers[positions[0]]).strip()}'")
            if spec.essential:
                result.essential_problems.append(f"'{spec.output}' conflicting columns ({detail})")
                result.findings.append(_finding("essential_conflict", category, spec.output, detail))
            else:
                result.findings.append(_finding("value_conflict", category, spec.output, detail))

    frame = pd.DataFrame(data, columns=columns)
    frame["_source_row_number"] = row_numbers
    result.frame = frame
    return result


def _sheet_signature(wb: openpyxl.Workbook, sheet_name: str, category: str) -> Tuple:
    rows = [tuple(r) for r in wb[sheet_name].iter_rows(values_only=True)]
    if not rows:
        return ((), 0)
    idx = _detect_header_row(rows, _accepted_names(catalogue.FIELDS_BY_CATEGORY[category]))
    header = tuple(normalize_header(h) for h in rows[idx])
    body = tuple(
        tuple(_safe_str(_norm_cell(v)) for v in r)
        for r in rows[idx + 1:]
        if not all(_is_blank(v) for v in r)
    )
    return (header, body)


def resolve_sheets(
    wb: openpyxl.Workbook,
    snapshot_date: str,
    configured: Optional[Dict[str, str]] = None,
) -> Tuple[Dict[str, Optional[str]], List[Dict[str, str]]]:
    """Pick the worksheet for each supported category in this workbook.

    Candidates are the configured name plus every catalogue name for the
    category, compared after normalization. Identical duplicates are
    harmless; materially different candidates are not guessed between.
    """
    configured = configured or {}
    findings: List[Dict[str, str]] = []
    chosen: Dict[str, Optional[str]] = {}
    by_norm: Dict[str, List[str]] = {}
    for name in wb.sheetnames:
        by_norm.setdefault(normalize_header(name), []).append(name)

    claimed: set = set()
    for category, spec in catalogue.SHEETS.items():
        wanted = [configured.get(category, "")] + list(spec.names)
        candidates: List[str] = []
        for w in wanted:
            for actual in by_norm.get(normalize_header(w), []):
                if w and actual not in candidates:
                    candidates.append(actual)
        claimed.update(candidates)
        if not candidates:
            chosen[category] = None
            if spec.expected_since is None or snapshot_date >= spec.expected_since:
                findings.append(_finding(
                    "missing_sheet", category, " / ".join(_unique(n for n in wanted if n))
                ))
            continue
        if len(candidates) > 1:
            signatures = {_sheet_signature(wb, c, category) for c in candidates}
            if len(signatures) > 1:
                chosen[category] = None
                findings.append(_finding(
                    "ambiguous_sheet", category, " / ".join(candidates),
                    "Tabs differ in columns or rows; remove one or ask the director "
                    "to update icms_catalogue.py.",
                ))
                continue
        chosen[category] = candidates[0]

    known = {normalize_header(n) for n in catalogue.KNOWN_EXCLUDED_SHEETS}
    for name in wb.sheetnames:
        if name not in claimed and normalize_header(name) not in known:
            findings.append(_finding("unfamiliar_sheet", "", name))
    return chosen, findings


# ── Frame builders ────────────────────────────────────────────────────────────

def _with_metadata(df: pd.DataFrame, snapshot_date: str, source_file_name: str) -> pd.DataFrame:
    out = df.copy()
    out.insert(0, "source_file_name", source_file_name)
    out.insert(0, "reporting_month", snapshot_date[:7])
    out.insert(0, "snapshot_date", snapshot_date)
    return out


def build_participant_frame(
    category_df: pd.DataFrame,
    snapshot_date: str,
    source_file_name: str,
) -> Tuple[pd.DataFrame, int]:
    """Add metadata and drop rows lacking a CHAMP ID or project.

    Returns ``(frame, dropped_row_count)``.
    """
    out = _with_metadata(category_df, snapshot_date, source_file_name)
    out = out.reindex(columns=METADATA_FIELDS + ALL_APPROVED_FIELDS + ["_source_row_number"])
    keep = pd.Series(True, index=out.index)
    for k in ROW_KEY_FIELDS:
        keep &= out[k].map(_safe_str).ne("")
    dropped = int((~keep).sum())
    return out.loc[keep].reset_index(drop=True), dropped


def build_vacant_slot_rows(
    vacant_df: pd.DataFrame,
    snapshot_date: str,
    source_file_name: str,
) -> pd.DataFrame:
    """Convert vacant slot records into flat-file rows (``is_vacant_slot = True``)."""
    if vacant_df.empty:
        return pd.DataFrame(columns=FLATFILE_COLUMNS)
    out = _with_metadata(_strip_internal_cols(vacant_df), snapshot_date, source_file_name)
    out["is_vacant_slot"] = True
    return out.reindex(columns=FLATFILE_COLUMNS)


# ── Duplicate detection ───────────────────────────────────────────────────────

def find_duplicates(df: pd.DataFrame) -> pd.DataFrame:
    """Return all rows sharing CHAMP ID + ICMS Project Name + ICMS Enrollment ID."""
    if any(k not in df.columns for k in KEY_FIELDS) or df.empty:
        return pd.DataFrame()
    keys = df[KEY_FIELDS].apply(lambda s: s.map(_safe_str))
    return df[keys.duplicated(keep=False)].copy()


def build_duplicate_warning_rows(
    dup_df: pd.DataFrame,
    run_timestamp: str,
    snapshot_date: str,
    source_file_name: str,
) -> pd.DataFrame:
    """Format duplicate rows into the warning-file schema."""
    warning_schema_cols = [
        "run_timestamp",
        "snapshot_date",
        "source_file_name",
        "CHAMP ID",
        "ICMS Project Name",
        "ICMS Enrollment ID",
        "ICMS Provider",
        "PH Building Name",
        "ICMS Case Manager",
        "source_row_number",
    ]
    rows = []
    for _, row in dup_df.iterrows():
        rows.append(
            {
                "run_timestamp": run_timestamp,
                "snapshot_date": snapshot_date,
                "source_file_name": source_file_name,
                "CHAMP ID": _safe_str(row.get("CHAMP ID")),
                "ICMS Project Name": _safe_str(row.get("ICMS Project Name")),
                "ICMS Enrollment ID": _safe_str(row.get("ICMS Enrollment ID")),
                "ICMS Provider": _safe_str(row.get("ICMS Provider")),
                "PH Building Name": _safe_str(row.get("PH Building Name")),
                "ICMS Case Manager": _safe_str(row.get("ICMS Case Manager")),
                "source_row_number": _safe_str(row.get("_source_row_number")),
            }
        )
    return pd.DataFrame(rows, columns=warning_schema_cols)


# ── Output file writers ───────────────────────────────────────────────────────

def update_dashboard_flatfile(
    df: pd.DataFrame,
    output_dir: Path,
    vacant_rows: pd.DataFrame,
) -> int:
    """Overwrite ``icms_current_dashboard_flatfile.xlsx``; return its row count."""
    path = output_dir / DASHBOARD_FILENAME
    out = _strip_internal_cols(df).copy()
    out["is_vacant_slot"] = False
    frames = [out.reindex(columns=FLATFILE_COLUMNS)]
    if vacant_rows is not None and not vacant_rows.empty:
        frames.append(vacant_rows.reindex(columns=FLATFILE_COLUMNS))
    records = [r for f in frames for r in f.to_dict("records")]
    combined = pd.DataFrame(records, columns=FLATFILE_COLUMNS)
    try:
        _write_excel(combined, path)
    except Exception as exc:
        raise RuntimeError(f"Could not write dashboard flat file: {exc}") from exc
    return len(combined)


def _sort_history(df: pd.DataFrame) -> pd.DataFrame:
    """Stable sort by snapshot date then source so import order never matters."""
    if df.empty or "snapshot_date" not in df.columns:
        return df.reset_index(drop=True)
    work = df.copy()
    work["_snap"] = pd.to_datetime(work["snapshot_date"], errors="coerce")
    work["_src"] = work.get("source_file_name", "")
    work = work.sort_values(["_snap", "_src"], kind="stable")
    return work.drop(columns=["_snap", "_src"]).reset_index(drop=True)


def update_snapshot_history(
    df: pd.DataFrame,
    snapshot_date: str,
    output_dir: Path,
) -> int:
    """Replace all rows for ``snapshot_date`` in ``icms_snapshot_history.csv``.

    Returns the total row count after the update.
    """
    path = output_dir / SNAPSHOT_HISTORY_FILENAME
    incoming = _stringify_frame(_strip_internal_cols(df))
    existing = _read_history_csv(path, "snapshot history")
    if not existing.empty and "snapshot_date" in existing.columns:
        existing = existing[existing["snapshot_date"] != snapshot_date]
    combined = pd.concat([existing, incoming], ignore_index=True, sort=False)
    combined = _sort_history(
        _order_columns(_stringify_frame(combined), METADATA_FIELDS + ALL_APPROVED_FIELDS)
    )
    try:
        combined.to_csv(path, index=False, encoding="utf-8")
    except Exception as exc:
        raise RuntimeError(f"Could not write snapshot history: {exc}") from exc
    return len(combined)


def update_exit_log_history(
    exit_df: pd.DataFrame,
    snapshot_date: str,
    source_file_name: str,
    output_dir: Path,
) -> int:
    """Replace this file's rows (same source name + snapshot date) in exit history.

    An empty ``exit_df`` (no exits, or the tab was not imported) leaves the
    stored history untouched. Returns the total row count after the update.
    """
    path = output_dir / EXIT_HISTORY_FILENAME
    existing = _read_history_csv(path, "exit log history")
    if exit_df is None or exit_df.empty:
        return len(existing)

    incoming = _with_metadata(
        _stringify_frame(_strip_internal_cols(exit_df)), snapshot_date, source_file_name
    )
    if not existing.empty and {"snapshot_date", "source_file_name"} <= set(existing.columns):
        existing = existing[
            ~(
                (existing["snapshot_date"] == snapshot_date)
                & (existing["source_file_name"] == source_file_name)
            )
        ]
    combined = pd.concat([existing, incoming], ignore_index=True, sort=False)
    combined = _sort_history(
        _order_columns(_stringify_frame(combined), METADATA_FIELDS + EXIT_HISTORY_FIELDS)
    )
    try:
        combined.to_csv(path, index=False, encoding="utf-8")
    except Exception as exc:
        raise RuntimeError(f"Could not write exit log history: {exc}") from exc
    return len(combined)


def _build_services_event_key_df(df: pd.DataFrame) -> pd.DataFrame:
    """Return deterministic service-event keys for deduplication/comparison.

    Rows with a Service ID are keyed on it. Others use the composite
    ``CHAMP|Enrollment|Service|Place|Service Date|Created Date``. A blank
    Service (as written by older releases that missed the "ICMS Service"
    rename) adopts the single non-blank Service seen elsewhere in ``df``
    for the same remaining components, so repaired and damaged copies of
    one event share a key.
    """
    work = pd.DataFrame(index=df.index)
    for col in SERVICES_LOG_FIELDS:
        src = df[col] if col in df.columns else pd.Series("", index=df.index)
        work[col] = src.fillna("").astype(str).str.strip()
    for col in ("Service Date", "Service Record Created Date"):
        parsed = _parse_date_series(work[col])
        work[col] = parsed.dt.strftime("%Y-%m-%d").where(parsed.notna(), work[col])
    for col in ("CHAMP ID", "ICMS Enrollment ID"):
        work[col] = _norm_id_series(work[col])

    service_id = work["Service ID"]
    has_service_id = service_id.ne("")
    base = (
        work["CHAMP ID"] + "|" + work["ICMS Enrollment ID"] + "|"
        + work["Place of Service"] + "|" + work["Service Date"] + "|"
        + work["Service Record Created Date"]
    )
    service = work["Service"]
    named = (~has_service_id) & service.ne("")
    names_per_base = (
        pd.DataFrame({"base": base[named], "service": service[named]})
        .drop_duplicates()
        .groupby("base")["service"]
    )
    unique_name = names_per_base.first()[names_per_base.nunique() == 1]
    effective_service = service.where(service.ne(""), base.map(unique_name).fillna(""))

    fallback = (
        work["CHAMP ID"] + "|" + work["ICMS Enrollment ID"] + "|"
        + effective_service + "|" + work["Place of Service"] + "|"
        + work["Service Date"] + "|" + work["Service Record Created Date"]
    )
    event_key = service_id.where(has_service_id, "__fallback__|" + fallback)
    return pd.DataFrame(
        {
            "service_event_key": event_key,
            "has_service_id": has_service_id.astype(bool),
        },
        index=df.index,
    )


def summarize_services_overlap(
    services_df: pd.DataFrame,
    output_dir: Path,
) -> Dict[str, float]:
    """Summarise incoming services-table overlap against cumulative history."""
    summary: Dict[str, float] = {
        "incoming_event_key_count": 0,
        "existing_overlap_rate": 0.0,
    }
    if services_df is None or services_df.empty:
        return summary
    existing = _read_history_csv(output_dir / SERVICES_HISTORY_FILENAME, "services log history")
    incoming = _stringify_frame(_strip_internal_cols(services_df))
    incoming["_side"] = "in"
    if not existing.empty:
        existing = existing.copy()
        existing["_side"] = "old"
    both = pd.concat([existing, incoming], ignore_index=True, sort=False).fillna("")
    keys = _build_services_event_key_df(both)["service_event_key"]
    incoming_keys = set(keys[both["_side"] == "in"])
    existing_keys = set(keys[both["_side"] == "old"])
    summary["incoming_event_key_count"] = len(incoming_keys)
    if incoming_keys and existing_keys:
        summary["existing_overlap_rate"] = len(incoming_keys & existing_keys) / len(incoming_keys)
    return summary


def update_services_log_history(
    services_df: pd.DataFrame,
    snapshot_date: str,
    source_file_name: str,
    output_dir: Path,
) -> Dict[str, float]:
    """Merge this file's Service Log rows into the deduplicated history.

    1. Rows previously imported from the same source name + snapshot date
       are removed first, so reprocessing a file replaces its old values
       (including a damaged representative row).
    2. Remaining rows are keyed by Service ID or the composite fallback.
    3. One row per key: per column, the earliest snapshot/source value that
       is non-blank. A blank written by a damaged import therefore never
       hides a value another snapshot supplied.

    Returns metrics used by GUI/logging.
    """
    path = output_dir / SERVICES_HISTORY_FILENAME
    metrics: Dict[str, float] = {
        "services_log_records_processed": 0,
        "services_log_history_total_rows": 0,
        "services_log_duplicate_records_dropped-diag": 0,
        "services_log_blank_service_id_count-diag": 0,
        "services_log_replaced_prior_rows-diag": 0,
    }
    existing = _read_history_csv(path, "services log history")
    if services_df is None or services_df.empty:
        metrics["services_log_history_total_rows"] = len(existing)
        return metrics

    incoming = _with_metadata(
        _stringify_frame(_strip_internal_cols(services_df)), snapshot_date, source_file_name
    )
    metrics["services_log_records_processed"] = int(len(incoming))
    metrics["services_log_blank_service_id_count-diag"] = int(
        incoming["Service ID"].fillna("").astype(str).str.strip().eq("").sum()
    )

    if not existing.empty and {"snapshot_date", "source_file_name"} <= set(existing.columns):
        same = (existing["snapshot_date"] == snapshot_date) & (
            existing["source_file_name"] == source_file_name
        )
        metrics["services_log_replaced_prior_rows-diag"] = int(same.sum())
        existing = existing[~same]

    combined = _stringify_frame(pd.concat([existing, incoming], ignore_index=True, sort=False))
    combined["_service_event_key"] = _build_services_event_key_df(combined)["service_event_key"]
    combined = _sort_history(combined)

    before = len(combined)
    data_cols = [c for c in combined.columns if c != "_service_event_key"]
    values = combined[data_cols].astype(object)
    values = values.mask(values.eq(""))
    values["_service_event_key"] = combined["_service_event_key"]
    deduped = (
        values.groupby("_service_event_key", sort=False)
        .first()
        .reset_index(drop=True)
        .reindex(columns=data_cols)
        .astype(object)
    )
    metrics["services_log_duplicate_records_dropped-diag"] = int(max(0, before - len(deduped)))

    ph_building = deduped["PH Building Name"].map(_safe_str)
    deduped.loc[ph_building.eq(""), "PH Building Name"] = "Scattered Site"
    deduped = _sort_history(
        _order_columns(_stringify_frame(deduped), METADATA_FIELDS + SERVICES_LOG_FIELDS)
    )
    try:
        deduped.to_csv(path, index=False, encoding="utf-8")
    except Exception as exc:
        raise RuntimeError(f"Could not write services log history: {exc}") from exc

    metrics["services_log_history_total_rows"] = int(len(deduped))
    return metrics


# ── Client/project/enrollment master ──────────────────────────────────────────

def _key_tuples(df: pd.DataFrame) -> List[Tuple[str, str, str]]:
    return list(zip(*(df[k].map(_safe_str) for k in KEY_FIELDS)))


def _prepare_keyed_history(df: pd.DataFrame, label: str) -> pd.DataFrame:
    missing = [k for k in KEY_FIELDS if k not in df.columns]
    if missing:
        raise RuntimeError(
            f"The stored {label} has no {', '.join(missing)} column, so the "
            "enrollment-level master cannot be built. Reprocess the source "
            "snapshots (batch run) to rebuild it."
        )
    out = df.copy()
    for k in ID_FIELDS:
        out[k] = _norm_id_series(out[k])
    out["ICMS Project Name"] = out["ICMS Project Name"].map(_safe_str)
    out["_snap"] = pd.to_datetime(out["snapshot_date"], errors="coerce")
    out["_src"] = out.get("source_file_name", "")
    return out


def _build_master_from_histories(output_dir: Path) -> pd.DataFrame:
    """Rebuild the master deterministically from the stored histories.

    Grain: CHAMP ID + ICMS Project Name + ICMS Enrollment ID. Participant
    values come from each enrollment's latest participant observation;
    exit details from its latest exit observation (latest snapshot, then
    latest exit date). Enrollments seen only in exit history are included.
    Status per enrollment: Active if in the latest participant snapshot,
    else Exited if a matching exit is recorded, else Unknown.
    """
    hist = _read_history_csv(output_dir / SNAPSHOT_HISTORY_FILENAME, "snapshot history")
    exit_hist = _read_history_csv(output_dir / EXIT_HISTORY_FILENAME, "exit log history")
    if hist.empty and exit_hist.empty:
        return pd.DataFrame(columns=MASTER_COLUMNS)

    latest_keys: set = set()
    if not hist.empty:
        hist = _prepare_keyed_history(hist, "snapshot history")
        hist = hist.sort_values(["_snap", "_src"], kind="stable")
        latest_participant = hist.groupby(KEY_FIELDS, sort=False).last().reset_index()
        latest_snap = hist["_snap"].max()
        if pd.notna(latest_snap):
            latest_keys = set(_key_tuples(hist[hist["_snap"] == latest_snap]))
    else:
        latest_participant = pd.DataFrame(columns=KEY_FIELDS)

    if not exit_hist.empty:
        exit_hist = _prepare_keyed_history(exit_hist, "exit log history")
        exit_hist["_event"] = pd.to_datetime(
            exit_hist.get("Check-Out Date", ""), errors="coerce"
        ).fillna(pd.to_datetime(exit_hist.get("Move-Out Date", ""), errors="coerce"))
        exit_hist = exit_hist.sort_values(["_snap", "_event", "_src"], kind="stable")
        latest_exit = exit_hist.groupby(KEY_FIELDS, sort=False).last().reset_index()
    else:
        latest_exit = pd.DataFrame(columns=KEY_FIELDS)

    master = latest_participant.reindex(columns=MASTER_COLUMNS, fill_value="").fillna("")
    exit_keys = set(_key_tuples(latest_exit)) if not latest_exit.empty else set()

    if not latest_exit.empty:
        exit_cols = [c for c in EXIT_FIELDS + ALL_APPROVED_FIELDS
                     if c in latest_exit.columns and c not in KEY_FIELDS]
        exit_side = latest_exit[KEY_FIELDS + exit_cols].rename(
            columns={c: f"{c}__exit" for c in exit_cols}
        )
        merged = master.merge(exit_side, on=KEY_FIELDS, how="left", indicator=True)
        matched = merged["_merge"] == "both"
        for col in exit_cols:
            exit_vals = merged[f"{col}__exit"].map(_safe_str)
            if col in EXIT_FIELDS:
                merged.loc[matched, col] = exit_vals[matched]
            else:
                fill = matched & merged[col].map(_safe_str).eq("")
                merged.loc[fill, col] = exit_vals[fill]
        master = merged[MASTER_COLUMNS].copy()

        participant_keys = set(_key_tuples(master)) if not master.empty else set()
        exit_only = latest_exit[
            [k not in participant_keys for k in _key_tuples(latest_exit)]
        ]
        if not exit_only.empty:
            extra = exit_only.reindex(columns=MASTER_COLUMNS, fill_value="").fillna("")
            master = pd.concat([master, extra], ignore_index=True)

    if master.empty:
        return master.reindex(columns=MASTER_COLUMNS)

    row_keys = _key_tuples(master)
    in_latest = np.array([k in latest_keys for k in row_keys], dtype=bool)
    in_exit = np.array([k in exit_keys for k in row_keys], dtype=bool)
    master["present_in_latest_snapshot"] = np.where(in_latest, "True", "False")
    master["client_status"] = np.select(
        [in_latest, in_exit], ["Active", "Exited"], default="Unknown"
    )
    master = master.sort_values(KEY_FIELDS, kind="stable")
    return _stringify_frame(master.reindex(columns=MASTER_COLUMNS, fill_value="")).reset_index(drop=True)


def compute_incoming_master_key_stats(
    df: pd.DataFrame,
    output_dir: Path,
) -> Tuple[int, int, int]:
    """Return ``(new, returning, not_present)`` enrollment counts vs the existing master."""
    incoming = set(_key_tuples(df)) if all(k in df.columns for k in KEY_FIELDS) else set()
    path = output_dir / MASTER_FILENAME
    existing: set = set()
    if path.exists():
        try:
            m = pd.read_excel(path, engine="openpyxl", dtype=str).fillna("")
            if all(k in m.columns for k in KEY_FIELDS):
                for k in ID_FIELDS:
                    m[k] = _norm_id_series(m[k])
                existing = set(_key_tuples(m))
        except Exception:
            existing = set()
    return len(incoming - existing), len(incoming & existing), len(existing - incoming)


def rebuild_client_project_master(output_dir: Path) -> Tuple[int, int, int]:
    """Rebuild and write the master; return ``(rows, exited, unknown)``."""
    path = output_dir / MASTER_FILENAME
    master = _build_master_from_histories(output_dir)
    status = master["client_status"] if not master.empty else pd.Series(dtype=str)
    try:
        _write_excel(master, path)
    except Exception as exc:
        raise RuntimeError(f"Could not write client/project master: {exc}") from exc
    return len(master), int((status == "Exited").sum()), int((status == "Unknown").sum())


def append_ingest_log(log_row: Dict, output_dir: Path) -> None:
    """Record a run in ``icms_ingest_log.csv`` (one row per snapshot date)."""
    path = output_dir / INGEST_LOG_FILENAME
    snap_date = str(log_row.get("snapshot_date", ""))
    existing = _read_history_csv(path, "ingest log")
    if not existing.empty and "snapshot_date" in existing.columns:
        existing = existing[existing["snapshot_date"] != snap_date]
    new_row = pd.DataFrame([{k: _safe_str(v) for k, v in log_row.items()}])
    combined = new_row if existing.empty else pd.concat([existing, new_row], ignore_index=True)
    try:
        combined.to_csv(path, index=False, encoding="utf-8")
    except Exception as exc:
        raise RuntimeError(f"Could not write ingest log: {exc}") from exc


def append_duplicate_warning(
    warning_df: pd.DataFrame,
    snapshot_date: str,
    output_dir: Path,
) -> None:
    """Replace this date's rows in ``icms_duplicate_client_project_warning.csv``."""
    path = output_dir / DUPLICATE_WARNING_FILENAME
    existing = _read_history_csv(path, "duplicate warning file")
    if not existing.empty and "snapshot_date" in existing.columns:
        existing = existing[existing["snapshot_date"] != snapshot_date]
    if warning_df.empty and existing.empty and not path.exists():
        return
    combined = _stringify_frame(pd.concat([existing, warning_df], ignore_index=True))
    try:
        combined.to_csv(path, index=False, encoding="utf-8")
    except Exception as exc:
        raise RuntimeError(f"Could not write duplicate warning file: {exc}") from exc


# ── Backups and per-file transactions ─────────────────────────────────────────

class OutputLockedError(RuntimeError):
    """Raised when an output file or folder cannot be written before a run."""

    def __init__(self, locked: Sequence[Path]):
        self.locked = [Path(p) for p in locked]
        super().__init__(format_output_lock_message(self.locked))


def format_output_lock_message(locked: Sequence[Path]) -> str:
    """User-facing explanation for locked or unwritable outputs."""
    # Prefer short names for files; keep the folder path when the dir itself failed.
    display: List[str] = []
    for p in locked:
        path = Path(p)
        display.append(path.name if path.suffix else str(path))
    listed = "\n".join(f"  - {n}" for n in display)
    return (
        "Cannot write to one or more output files.\n\n"
        f"{listed}\n\n"
        "Close these files in Excel or another program, wait for OneDrive sync "
        "to finish, then try again."
    )


def _probe_path_writable(path: Path) -> bool:
    """Return True if ``path`` can be opened for update (Excel locks fail here)."""
    try:
        with open(path, "r+b"):
            pass
        return True
    except PermissionError:
        return False
    except OSError as exc:
        # WinError 32: file in use; errno 13: EACCES
        if getattr(exc, "winerror", None) == 32 or getattr(exc, "errno", None) in (13, 11):
            return False
        raise


def check_outputs_writable(output_dir: Path) -> List[Path]:
    """Probe existing outputs and the output folder for write access.

    Returns locked/unwritable paths (empty list if everything is writable).
    Does not create missing output files. The output directory itself is
    included in the result when a write probe there fails.
    """
    output_dir = Path(output_dir)
    locked: List[Path] = []

    probe = output_dir / ".icms_write_probe"
    try:
        with open(probe, "wb") as fh:
            fh.write(b"0")
        probe.unlink(missing_ok=True)
    except OSError:
        probe.unlink(missing_ok=True)
        locked.append(output_dir)
        return locked

    for name in OUTPUT_FILENAMES:
        path = output_dir / name
        if path.exists() and not _probe_path_writable(path):
            locked.append(path)
    return locked


def ensure_outputs_writable(output_dir: Path) -> None:
    """Raise ``OutputLockedError`` if any existing output cannot be written."""
    locked = check_outputs_writable(output_dir)
    if locked:
        raise OutputLockedError(locked)


def existing_outputs(output_dir: Path) -> List[Path]:
    return [output_dir / n for n in OUTPUT_FILENAMES if (output_dir / n).exists()]


def output_format_is_current(output_dir: Path) -> bool:
    marker = output_dir / FORMAT_MARKER_FILENAME
    if not existing_outputs(output_dir):
        return True
    try:
        return marker.read_text(encoding="utf-8").strip() == OUTPUT_FORMAT_VERSION
    except OSError:
        return False


def write_format_marker(output_dir: Path) -> None:
    (output_dir / FORMAT_MARKER_FILENAME).write_text(OUTPUT_FORMAT_VERSION, encoding="utf-8")


def backup_outputs(output_dir: Path, reason: str) -> Optional[Path]:
    """Copy existing output files to ``backups/<timestamp> <reason>/``."""
    files = existing_outputs(output_dir)
    if not files:
        return None
    stamp = datetime.now().strftime("%Y-%m-%d %H%M%S")
    dest = output_dir / BACKUP_DIRNAME / f"{stamp} {reason}"
    dest.mkdir(parents=True, exist_ok=True)
    for f in files:
        shutil.copy2(f, dest / f.name)
    return dest


class _OutputTransaction:
    """Restore every output file if a file's writes fail part-way."""

    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self._tmp: Optional[Path] = None
        self._present: List[str] = []

    def __enter__(self):
        self._tmp = Path(tempfile.mkdtemp(prefix="icms_txn_"))
        for f in existing_outputs(self.output_dir):
            shutil.copy2(f, self._tmp / f.name)
            self._present.append(f.name)
        return self

    def __exit__(self, exc_type, exc, tb):
        try:
            if exc_type is not None:
                for name in OUTPUT_FILENAMES:
                    target = self.output_dir / name
                    if name in self._present:
                        shutil.copy2(self._tmp / name, target)
                    elif target.exists():
                        target.unlink()
        finally:
            shutil.rmtree(self._tmp, ignore_errors=True)
        return False


# ── Pipeline ──────────────────────────────────────────────────────────────────

STEP_LABELS: List[str] = [
    "Opening workbook",
    "Reading sheets",
    "Checking structure",
    "Updating snapshot and exit history",
    "Updating services history",
    "Updating client/project master",
    "Writing dashboard and logs",
]

LogFn = Callable[..., None]


def _default_log(msg: str, level: str = "info") -> None:
    encoding = getattr(sys.stdout, "encoding", None) or "utf-8"
    print(msg.encode(encoding, "replace").decode(encoding))


class FileImportError(Exception):
    """A file could not be imported; nothing from it was written."""

    def __init__(self, message: str, findings: Optional[List[Dict[str, str]]] = None):
        super().__init__(message)
        self.findings = findings or []


@dataclass
class FileSpec:
    path: Path
    snapshot_date: str
    date_source: str = ""
    generated_on: Optional[str] = None
    filename_date: Optional[str] = None


def process_snapshot_file(
    spec: FileSpec,
    output_dir: Path,
    *,
    password: Optional[str] = "",
    sheet_names: Optional[Dict[str, str]] = None,
    publish_dashboard: bool = True,
    defer_master_rebuild: bool = False,
    log: LogFn = _default_log,
    step: Callable[[int], None] = lambda n: None,
) -> Dict:
    """Import one snapshot workbook; return a result dict.

    Reads and validates everything before writing. All writes for the file
    happen inside a transaction: if any write fails, every output file is
    restored to its prior state. A file that fails is never published.
    """
    input_path = Path(spec.path)
    snapshot_date = spec.snapshot_date
    source_file_name = input_path.name
    run_ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    findings: List[Dict[str, str]] = []
    notes: List[str] = []
    base = {
        "source_file_name": source_file_name,
        "snapshot_date": snapshot_date,
        "snapshot_date_source": spec.date_source,
        "generated_on_date": spec.generated_on or "",
        "filename_date": spec.filename_date or "",
    }

    try:
        # ── 1. Open ────────────────────────────────────────────────────────
        step(1)
        log(f"Opening workbook: {source_file_name}")
        wb, was_encrypted = open_workbook(input_path, password=password)
        if was_encrypted:
            log("File was password-protected — decrypted in memory")

        try:
            generated = read_generated_on(wb)
            if generated and not spec.generated_on:
                base["generated_on_date"] = generated
            if not generated:
                findings.append(_finding("generated_on_missing", "", catalogue.METADATA_SHEET))
            from_name = spec.filename_date or parse_date_from_filename(source_file_name)
            if generated and (generated != snapshot_date or (from_name and from_name != generated)):
                findings.append(_finding(
                    "date_mismatch", "", source_file_name,
                    f"Imported as {snapshot_date}; Generated On {generated}; "
                    f"filename {from_name or 'none'}",
                ))
            log(f"Snapshot date: {snapshot_date}"
                + (f" (from {spec.date_source})" if spec.date_source else ""))

            # ── 2. Resolve and read sheets ─────────────────────────────────
            step(2)
            chosen, sheet_findings = resolve_sheets(wb, snapshot_date, sheet_names)
            findings.extend(sheet_findings)
            if not chosen["participant"]:
                raise FileImportError(
                    "Participant Level Summary tab not found (or more than one "
                    f"differing candidate). Sheets: {', '.join(wb.sheetnames)}",
                    findings,
                )
            reads: Dict[str, CategoryRead] = {}
            for category, sheet in chosen.items():
                if sheet:
                    log(f"Reading sheet: {sheet}")
                    reads[category] = read_category(wb, sheet, category, snapshot_date)
                    findings.extend(reads[category].findings)
        finally:
            wb.close()

        # ── 3. Structure checks ────────────────────────────────────────────
        step(3)
        participant = reads["participant"]
        if participant.essential_problems:
            raise FileImportError(
                "Participant Level Summary cannot be imported: "
                + "; ".join(participant.essential_problems),
                findings,
            )
        skipped: List[str] = []
        frames: Dict[str, pd.DataFrame] = {}
        for category in ("vacant", "exit", "services"):
            label = catalogue.CATEGORY_LABELS[category]
            read = reads.get(category)
            if read is None:
                frames[category] = pd.DataFrame()
                if any(f["kind"] in ("missing_sheet", "ambiguous_sheet")
                       and f["category"] == label for f in findings):
                    skipped.append(label)
                continue
            if read.essential_problems:
                skipped.append(label)
                findings.append(_finding(
                    "category_skipped", category, label, "; ".join(read.essential_problems)
                ))
                frames[category] = pd.DataFrame()
                continue
            frames[category] = read.frame
        for label in skipped:
            log(f"WARNING: {label} not imported for this file — see structural findings.",
                level="warn")
            notes.append(f"{label} not imported")

        df, dropped_rows = build_participant_frame(
            participant.frame, snapshot_date, source_file_name
        )
        if df.empty:
            blank = [k for k in ROW_KEY_FIELDS
                     if participant.frame[k].map(_safe_str).eq("").all()]
            detail = (f"{' and '.join(blank)} is blank on every row"
                      if blank else "no row has both a CHAMP ID and an ICMS Project Name")
            if blank:
                for k in blank:
                    findings.append(_finding("essential_missing", "participant", k,
                                             "column present but blank on every row"))
            raise FileImportError(f"Participant Level Summary cannot be imported: {detail}.",
                                  findings)
        if dropped_rows:
            notes.append(f"Dropped {dropped_rows} participant row(s) with blank CHAMP ID or project")
            log(f"WARNING: dropped {dropped_rows} row(s) with blank CHAMP ID or project",
                level="warn")
        blank_enrollment = int(df["ICMS Enrollment ID"].map(_safe_str).eq("").sum())
        if blank_enrollment:
            findings.append(_finding(
                "blank_enrollment", "participant", "ICMS Enrollment ID",
                f"{blank_enrollment} row(s); kept as separate blank-enrollment rows",
            ))

        exit_df = frames["exit"]
        if not exit_df.empty:
            keep = pd.Series(True, index=exit_df.index)
            for k in ROW_KEY_FIELDS:
                keep &= exit_df[k].map(_safe_str).ne("")
            exit_df = exit_df.loc[keep]
        services_df = frames["services"]
        if not services_df.empty:
            services_df = services_df[services_df["CHAMP ID"].map(_safe_str).ne("")]
        vacant_rows = build_vacant_slot_rows(frames["vacant"], snapshot_date, source_file_name)

        log(f"Participant rows: {len(df)}")
        log(f"Vacant slot records: {len(vacant_rows)}")
        log(f"Exit log records: {len(exit_df)}")
        log(f"Services log records: {len(services_df)}")

        dup_df = find_duplicates(df)
        dup_count = len(dup_df)
        if dup_count:
            n_groups = dup_df.groupby(KEY_FIELDS, sort=False).ngroups
            log(f"WARNING: {n_groups} enrollment(s) appear on more than one row "
                f"({dup_count} rows)", level="warn")
            notes.append(f"Duplicate enrollment rows: {n_groups}")

        new_rows, returning_rows, not_present = compute_incoming_master_key_stats(df, output_dir)
        services_overlap = summarize_services_overlap(services_df, output_dir)

        # ── 4–7. Writes (all-or-nothing for this file) ─────────────────────
        with _OutputTransaction(output_dir):
            step(4)
            history_count = update_snapshot_history(df, snapshot_date, output_dir)
            log(f"Snapshot history: {history_count} total rows")
            exit_history_count = update_exit_log_history(
                exit_df, snapshot_date, source_file_name, output_dir
            )
            log(f"Exit log history: {exit_history_count} total rows")

            step(5)
            services_metrics = update_services_log_history(
                services_df, snapshot_date, source_file_name, output_dir
            )
            log("Services log history: "
                f"{services_metrics['services_log_history_total_rows']} total rows "
                f"({services_metrics['services_log_duplicate_records_dropped-diag']} "
                "duplicate rows merged)")

            step(6)
            master_total = exited_count = unknown_count = 0
            if defer_master_rebuild:
                log("Master rebuild deferred until the batch completes")
            else:
                master_total, exited_count, unknown_count = rebuild_client_project_master(output_dir)
                log(f"Master: {master_total} enrollments ({exited_count} exited, "
                    f"{unknown_count} unknown)")

            step(7)
            dashboard_count = 0
            if publish_dashboard:
                dashboard_count = update_dashboard_flatfile(df, output_dir, vacant_rows)
                log(f"Dashboard flat file: {dashboard_count} rows "
                    f"({len(vacant_rows)} vacant slot rows)")
            append_duplicate_warning(
                build_duplicate_warning_rows(dup_df, run_ts, snapshot_date, source_file_name),
                snapshot_date, output_dir,
            )
            file_groups = group_findings([(source_file_name, findings)])
            append_ingest_log(
                {
                    "run_timestamp": run_ts,
                    **base,
                    "sheet_name": chosen["participant"] or "",
                    "vacant_sheet_name": chosen.get("vacant") or "",
                    "exit_sheet_name": chosen.get("exit") or "",
                    "services_sheet_name": chosen.get("services") or "",
                    "input_row_count": len(participant.frame),
                    "vacant_slot_rows_injected": len(vacant_rows),
                    "exit_log_records_processed": len(exit_df),
                    "current_dashboard_row_count": dashboard_count,
                    "dashboard_published": bool(publish_dashboard),
                    "history_total_row_count_after_update": history_count,
                    "exit_history_total_row_count_after_update": exit_history_count,
                    "master_row_count_after_update": master_total,
                    "new_client_project_rows": new_rows,
                    "returning_client_project_rows": returning_rows,
                    "not_present_in_latest_snapshot_count": not_present,
                    "master_exited_count": exited_count,
                    "master_unknown_count": unknown_count,
                    "duplicate_client_project_count": dup_count,
                    "services_log_records_processed": services_metrics["services_log_records_processed"],
                    "services_log_history_total_rows": services_metrics["services_log_history_total_rows"],
                    "services_log_duplicate_records_dropped-diag":
                        services_metrics["services_log_duplicate_records_dropped-diag"],
                    "services_log_replaced_prior_rows-diag":
                        services_metrics["services_log_replaced_prior_rows-diag"],
                    "services_log_blank_service_id_count-diag":
                        services_metrics["services_log_blank_service_id_count-diag"],
                    "services_log_existing_key_overlap_rate-diag":
                        round(services_overlap["existing_overlap_rate"], 4),
                    "header_row_detected-diag": participant.header_row,
                    "dropped_blank_key_rows-diag": dropped_rows,
                    "blank_enrollment_id_rows-diag": blank_enrollment,
                    "skipped_categories": "; ".join(skipped),
                    "structural_findings": format_findings_text(file_groups).replace("\n", " | "),
                    "password_protected_input": was_encrypted,
                    "run_status": "success",
                    "notes": "; ".join(notes),
                },
                output_dir,
            )
        log("✓  File imported")
        return {
            **base,
            "success": True,
            "sheet_names": dict(chosen),
            "input_row_count": len(participant.frame),
            "participant_count": len(df),
            "dashboard_count": dashboard_count,
            "dashboard_published": bool(publish_dashboard),
            "vacant_slot_count": len(vacant_rows),
            "exit_log_count": len(exit_df),
            "services_log_count": services_metrics["services_log_records_processed"],
            "services_log_history_count": services_metrics["services_log_history_total_rows"],
            "services_log_dupes_dropped": services_metrics["services_log_duplicate_records_dropped-diag"],
            "history_count": history_count,
            "exit_history_count": exit_history_count,
            "master_total": master_total,
            "exited_count": exited_count,
            "unknown_count": unknown_count,
            "dup_flag_count": dup_count,
            "skipped_categories": skipped,
            "findings": findings,
            "notes": "; ".join(notes),
            "was_encrypted": was_encrypted,
            "_frames": (df, vacant_rows),
        }

    except Exception as exc:
        message = str(exc).strip() or exc.__class__.__name__
        if isinstance(exc, FileImportError):
            findings = exc.findings
        log(f"ERROR: {message}", level="error")
        log("Nothing from this file was written; existing outputs are unchanged.",
            level="error")
        try:
            append_ingest_log(
                {
                    "run_timestamp": run_ts,
                    **base,
                    "run_status": "failed",
                    "notes": message,
                    "structural_findings": format_findings_text(
                        group_findings([(source_file_name, findings)])
                    ).replace("\n", " | "),
                },
                output_dir,
            )
        except Exception:
            pass
        return {**base, "success": False, "error": message, "findings": findings}


def latest_history_snapshot_date(output_dir: Path) -> Optional[str]:
    hist_path = output_dir / SNAPSHOT_HISTORY_FILENAME
    if not hist_path.exists():
        return None
    try:
        dates = pd.read_csv(hist_path, dtype=str, usecols=["snapshot_date"])["snapshot_date"]
    except Exception:
        return None
    dates = dates.dropna()
    return max(dates) if len(dates) else None


def run_snapshot_files(
    files: Sequence[FileSpec],
    output_dir: Path,
    *,
    password: Optional[str] = "",
    sheet_names: Optional[Dict[str, str]] = None,
    publish_dashboard: bool = True,
    reprocessing: bool = False,
    log: LogFn = _default_log,
    step: Callable[[int], None] = lambda n: None,
) -> Dict:
    """Import one or more snapshots; the single entry point for GUI and CLI.

    - Files run oldest first (snapshot date, then filename), so selection
      order never changes the result.
    - Existing outputs are backed up automatically before a run that
      reprocesses already-imported files, and before the first run that
      upgrades outputs written by an earlier release.
    - The master is rebuilt once from the full histories at the end.
    - With ``publish_dashboard``, a single file publishes itself; a batch
      publishes only its newest successful file, and only if that file is not
      older than a snapshot already in history.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    ensure_outputs_writable(output_dir)
    ordered = sorted(files, key=lambda f: (f.snapshot_date, Path(f.path).name))
    is_batch = len(ordered) > 1
    prior_latest = latest_history_snapshot_date(output_dir)

    backup_dir: Optional[Path] = None
    if existing_outputs(output_dir) and (reprocessing or not output_format_is_current(output_dir)):
        reason = "before reprocessing" if reprocessing else "before format upgrade"
        backup_dir = backup_outputs(output_dir, reason)
        if backup_dir:
            log(f"Backed up existing outputs to {backup_dir}")

    results: List[Dict] = []
    for idx, spec in enumerate(ordered):
        if is_batch:
            log(f"Batch file {idx + 1} of {len(ordered)}: {Path(spec.path).name} "
                f"({spec.snapshot_date})")
        results.append(process_snapshot_file(
            spec, output_dir,
            password=password,
            sheet_names=sheet_names,
            publish_dashboard=publish_dashboard and not is_batch,
            defer_master_rebuild=is_batch,
            log=log, step=step,
        ))

    succeeded = [r for r in results if r["success"]]
    master_total = exited = unknown = 0
    master_error = ""
    if succeeded:
        if is_batch:
            log("Rebuilding client/project master from full history…")
            try:
                master_total, exited, unknown = rebuild_client_project_master(output_dir)
                log(f"Master: {master_total} enrollments ({exited} exited, {unknown} unknown)")
            except Exception as exc:
                master_error = str(exc).strip() or exc.__class__.__name__
                log(f"ERROR: master rebuild failed: {master_error}", level="error")
        else:
            r = succeeded[0]
            master_total, exited, unknown = r["master_total"], r["exited_count"], r["unknown_count"]
        write_format_marker(output_dir)

    dashboard_source = ""
    dashboard_count = 0
    dashboard_note = ""
    if not publish_dashboard:
        dashboard_note = "Dashboard publish was turned off for this run."
    elif succeeded and is_batch:
        newest = succeeded[-1]
        if prior_latest and newest["snapshot_date"] < prior_latest:
            dashboard_note = (
                f"Dashboard not updated: newest file in this batch "
                f"({newest['snapshot_date']}) is older than a snapshot already "
                f"imported ({prior_latest})."
            )
        else:
            df, vacant_rows = newest["_frames"]
            dashboard_count = update_dashboard_flatfile(df, output_dir, vacant_rows)
            dashboard_source = newest["source_file_name"]
            newest["dashboard_published"] = True
            newest["dashboard_count"] = dashboard_count
            log(f"Dashboard flat file published from {dashboard_source}: {dashboard_count} rows")
    elif succeeded:
        dashboard_source = succeeded[0]["source_file_name"]
        dashboard_count = succeeded[0]["dashboard_count"]
    if dashboard_note:
        log(dashboard_note, level="warn")

    for r in results:
        r.pop("_frames", None)

    groups = group_findings([(r["source_file_name"], r.get("findings", [])) for r in results])
    findings_text = format_findings_text(groups, total_files=len(results))
    skipped = sorted({c for r in succeeded for c in r.get("skipped_categories", [])})
    if findings_text:
        log("Structural findings:\n" + findings_text, level="warn")

    failed = [r for r in results if not r["success"]]
    return {
        "success": bool(succeeded) and not master_error,
        "is_batch": is_batch,
        "batch_total": len(results),
        "batch_success": len(succeeded),
        "batch_failed": len(failed),
        "batch_results": results,
        "batch_failed_files": [
            {"source_file_name": r["source_file_name"], "snapshot_date": r["snapshot_date"],
             "error": r.get("error", "")}
            for r in failed
        ],
        "dashboard_published": bool(dashboard_source),
        "dashboard_source": dashboard_source,
        "dashboard_count": dashboard_count,
        "dashboard_note": dashboard_note,
        "master_total": master_total,
        "master_exited_count": exited,
        "master_unknown_count": unknown,
        "master_error": master_error,
        "skipped_categories": skipped,
        "finding_groups": groups,
        "findings_text": findings_text,
        "backup_dir": str(backup_dir) if backup_dir else "",
        "error": master_error or (failed[0].get("error", "") if failed and not succeeded else ""),
        "output_dir": str(output_dir.resolve()),
    }


# ── Already-imported checks (used before a run to offer reprocessing) ─────────

def load_integrated_file_keys(output_dir: Path) -> set:
    """Return ``{(source_file_name, snapshot_date)}`` already imported."""
    keys: set = set()
    hist = output_dir / SNAPSHOT_HISTORY_FILENAME
    if hist.exists():
        try:
            h = pd.read_csv(hist, dtype=str, usecols=["source_file_name", "snapshot_date"]).fillna("")
            keys |= {(s.strip(), d.strip()) for s, d in h.drop_duplicates().itertuples(index=False)
                     if s.strip() and d.strip()}
        except Exception:
            pass
    ingest = output_dir / INGEST_LOG_FILENAME
    if ingest.exists():
        try:
            g = pd.read_csv(ingest, dtype=str).fillna("")
            if {"source_file_name", "snapshot_date", "run_status"} <= set(g.columns):
                ok = g[g["run_status"].str.lower() == "success"]
                keys |= {(s.strip(), d.strip()) for s, d in
                         ok[["source_file_name", "snapshot_date"]].itertuples(index=False)
                         if s.strip() and d.strip()}
        except Exception:
            pass
    return keys


def already_imported(files: Sequence[FileSpec], output_dir: Path) -> List[FileSpec]:
    """Files whose (name, date) already appear in this output folder."""
    keys = load_integrated_file_keys(output_dir)
    return [f for f in files if (Path(f.path).name.strip(), f.snapshot_date.strip()) in keys]


# ── Compatibility helpers for diagnostics/ scripts ────────────────────────────
# The diagnostics predate the catalogue; these keep them runnable.

def load_aliases(script_dir: Path) -> Dict[str, str]:
    """Retired: mappings now live in icms_catalogue.py."""
    return {}


def find_header_row(wb, sheet_name: str, candidate_fields: List[str], max_scan: int = 5) -> int:
    rows = [tuple(r) for r in wb[sheet_name].iter_rows(values_only=True)]
    return _detect_header_row(rows, {normalize_header(c) for c in candidate_fields}, max_scan)


def read_sheet(wb, sheet_name: str, header_row: int = 0, **_ignored) -> pd.DataFrame:
    """Raw sheet read with the header names as-is (diagnostics only)."""
    rows = [tuple(r) for r in wb[sheet_name].iter_rows(values_only=True)]
    if not rows:
        return pd.DataFrame()
    headers = [str(h).strip() if h is not None and str(h).strip() else f"_unnamed_col_{i}"
               for i, h in enumerate(rows[header_row])]
    df = pd.DataFrame(rows[header_row + 1:], columns=headers)
    df["_source_row_number"] = range(header_row + 2, header_row + 2 + len(df))
    for col in headers:
        df[col] = df[col].map(normalize_identifier if col in ID_FIELDS else _norm_cell)
    return df.dropna(subset=headers, how="all").reset_index(drop=True)


def _read_default(wb, category: str, snapshot_date: str = "9999-12-31") -> pd.DataFrame:
    chosen, _ = resolve_sheets(wb, snapshot_date)
    if not chosen.get(category):
        return pd.DataFrame()
    return read_category(wb, chosen[category], category, snapshot_date).frame


def read_vacant_slots(wb) -> pd.DataFrame:
    return _read_default(wb, "vacant")


def read_exit_log(wb) -> pd.DataFrame:
    return _read_default(wb, "exit")


def read_services_log(wb, sheet_name: str = "") -> pd.DataFrame:
    return _read_default(wb, "services")


def resolve_fields(df_columns: List[str], aliases: Optional[Dict[str, str]] = None):
    """Map participant outputs to raw headers via the catalogue.

    Returns ``(resolved, unresolved_essential, unresolved_other)``.
    """
    by_norm = {normalize_header(c): c for c in df_columns}
    resolved: Dict[str, str] = {}
    missing_essential: List[str] = []
    missing_other: List[str] = []
    for spec in catalogue.PARTICIPANT_FIELDS:
        actual = next((by_norm[normalize_header(n)] for n in spec.accepted
                       if normalize_header(n) in by_norm), None)
        if actual:
            resolved[spec.output] = actual
        elif spec.essential:
            missing_essential.append(spec.output)
        else:
            missing_other.append(spec.output)
    return resolved, missing_essential, missing_other


def reduce_to_approved_fields(df, resolved, snapshot_date, source_file_name, **_ignored):
    cols = {canonical: df[actual] for canonical, actual in resolved.items()}
    frame = pd.DataFrame(cols, index=df.index).reindex(columns=ALL_APPROVED_FIELDS)
    if "_source_row_number" in df.columns:
        frame["_source_row_number"] = df["_source_row_number"]
    return build_participant_frame(frame, snapshot_date, source_file_name)[0]


def record_master_inputs(exit_df, snapshot_date, source_file_name, output_dir) -> None:
    update_exit_log_history(exit_df, snapshot_date, source_file_name, output_dir)


def update_client_project_master(df, snapshot_date, source_file_name, output_dir, exit_df):
    record_master_inputs(exit_df, snapshot_date, source_file_name, output_dir)
    new, returning, absent = compute_incoming_master_key_stats(df, output_dir)
    total, exited, unknown = rebuild_client_project_master(output_dir)
    return total, new, returning, absent, exited, unknown
